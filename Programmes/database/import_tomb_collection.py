"""Import the reviewed historical Tomb collection without online rechecking."""

import hashlib
import json
import os
from pathlib import Path
import shutil
import sqlite3
import tempfile

import openpyxl
import pandas as pd

import database_initialize as loader


ROOT = Path(__file__).resolve().parents[2]
EXCEL_PATH = ROOT / "Programmes/database/data.xlsx"
DB_PATH = ROOT / "Programmes/database/data.db"
FULL_PATH = ROOT / "Programmes/datacollection/core_data/full_collection_neural.xlsx"
SOURCE_PATH = ROOT / (
    "Programmes/datacollection/runs/"
    "strict3333_transcript_splicing_v2_dedup_access_20260711_192604/"
    "outputs/08_agent_merged.xlsx"
)
CLASSIFICATION_PATH = ROOT / "Programmes/datacollection/archive/full_collection_12072026.xlsx"

# Match the old category meanings, rather than reusing changed category numbers.
CLASSIFICATION_MAP = {
    "Transcript and Isoform Model Resources": "I_1",
    "Transcript/Isoform Expression and Usage Resources": "I_2",
    "Alternative Splicing Event Resources": "II_1",
    "Splice-Site and Junction Resources": "II_2",
    "Splicing-Regulator Resources": "III_1",
    "Splicing-Associated Variant Resources": "III_2",
    "circRNA Resources": "IV_1",
    "Fusion/Chimeric RNA Resources": "IV_2",
    "lncRNA Isoform Resources": "IV_3",
    "NMD/Cryptic Transcript Resources": "IV_5",
}


def read_rows(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        values = book.worksheets[0].iter_rows(values_only=True)
        columns = [item["column_name"] for item in loader.parse_columns(next(values))]
        return [dict(zip(columns, row)) for row in values if any(v is not None for v in row)]
    finally:
        book.close()


def row_id(row):
    return str(row["id"]).strip().removesuffix(".0")


def build_tomb_records(headers, full_rows):
    full_ids = {row_id(row) for row in full_rows}
    full_names = {row["database_name"].strip().casefold() for row in full_rows}
    candidates = [
        row for row in read_rows(SOURCE_PATH)
        if row["db_type_confirmation"] == "yes" and row["accessibility"] == "dead"
        and row_id(row) not in full_ids
        and row["database_name"].strip().casefold() not in full_names
    ]
    if len(candidates) != 95 or len({row_id(row) for row in candidates}) != 95:
        raise ValueError("The historical candidate set has changed; expected 95 unique databases.")

    old_classification = {row_id(row): row for row in read_rows(CLASSIFICATION_PATH)}
    metadata = loader.parse_columns(headers)
    loader.validate_schema(metadata)
    records = []
    for candidate in candidates:
        record = {}
        for column in metadata:
            name = column["column_name"]
            value = candidate.get(name)
            if value is None or value == "":
                value = None if column["data_type"] in {
                    loader.TAG_NUMERIC, loader.TAG_NUMERIC_CITE,
                    loader.TAG_BOOL, loader.TAG_BOOL_ACCESS,
                    loader.TAG_WORD_URL, loader.TAG_WORD_DOI,
                } else "unknown"
            record[name] = value
        old_names = old_classification[row_id(candidate)].get("classification_sub", "")
        codes = dict.fromkeys(
            CLASSIFICATION_MAP[name.strip()]
            for name in str(old_names).split(";")
            if name.strip() in CLASSIFICATION_MAP
        )
        record.update(
            accessibility=False, main_collection="no", tomb=1,
            classification_code=";".join(codes) or "unknown",
        )
        records.append(record)

    frame = pd.DataFrame(records, columns=[item["column_name"] for item in metadata])
    raw = frame.rename(columns={item["column_name"]: item["original_name"] for item in metadata})
    loader.normalize_loader_values(raw, frame, metadata)
    loader.validate_data(metadata, frame)
    return [
        {name: None if pd.isna(value) else value for name, value in zip(frame.columns, values)}
        for values in frame.itertuples(index=False, name=None)
    ]


def main():
    excel_before = EXCEL_PATH.read_bytes()
    book = openpyxl.load_workbook(EXCEL_PATH)
    sheet = book.worksheets[0]
    headers = [cell.value for cell in sheet[1]]
    if "tomb" not in headers:
        headers.append("tomb")
    columns = [item["column_name"] for item in loader.parse_columns(headers)]
    workbook_rows = read_rows(EXCEL_PATH)
    full_rows = read_rows(FULL_PATH)
    records = build_tomb_records(headers, full_rows)

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    replaced_workbook = False
    committed = False
    try:
        db_rows = [dict(row) for row in conn.execute("SELECT * FROM database_info ORDER BY id")]
        workbook_by_id = {row_id(row): row for row in workbook_rows}
        db_by_id = {row_id(row): row for row in db_rows}
        if len(workbook_by_id) != len(workbook_rows) or len(db_by_id) != len(db_rows):
            raise ValueError("Existing records have duplicate IDs.")
        if workbook_by_id.keys() != db_by_id.keys():
            raise ValueError("Excel and SQLite membership differs; nothing was changed.")
        for key, row in workbook_by_id.items():
            if int(row.get("tomb") or 0) != int(db_by_id[key].get("tomb") or 0):
                raise ValueError(f"Excel and SQLite Tomb membership differs for ID {key}.")
        pending = [row for row in records if row_id(row) not in db_by_id]
        for row in records:
            if row_id(row) in db_by_id and not db_by_id[row_id(row)].get("tomb"):
                raise ValueError(f"Candidate {row_id(row)} already belongs to a current collection.")
        if not pending and "tomb" in db_rows[0]:
            print(json.dumps({"added": 0, "tomb": len(records), "status": "already imported"}))
            return

        tomb_column = columns.index("tomb") + 1
        sheet.cell(1, tomb_column, "tomb")
        for index, row in enumerate(workbook_rows, start=2):
            sheet.cell(index, tomb_column, int(row.get("tomb") or 0))
        for row in pending:
            sheet.append([row[name] for name in columns])
            if row["database_url"]:
                cell = sheet.cell(sheet.max_row, columns.index("database_url") + 1)
                cell.hyperlink = row["database_url"]

        backup_root = ROOT / "PreviousData"
        backup_root.mkdir(exist_ok=True)
        backup = Path(tempfile.mkdtemp(prefix="tomb_import_", dir=backup_root))
        (backup / "data.xlsx").write_bytes(excel_before)
        backup_conn = sqlite3.connect(backup / "data.db")
        try:
            conn.backup(backup_conn)
        finally:
            backup_conn.close()
        staged = backup / "data.with_tomb.xlsx"
        book.save(staged)
        book.close()

        conn.execute("BEGIN IMMEDIATE")
        if [dict(row) for row in conn.execute("SELECT * FROM database_info ORDER BY id")] != db_rows:
            raise ValueError("SQLite changed during preparation; import cancelled.")
        if EXCEL_PATH.read_bytes() != excel_before:
            raise ValueError("Excel changed during preparation; import cancelled.")
        db_columns = {row[1] for row in conn.execute("PRAGMA table_info(database_info)")}
        if "tomb" not in db_columns:
            conn.execute("ALTER TABLE database_info ADD COLUMN tomb INTEGER NOT NULL DEFAULT 0")
        if not conn.execute("SELECT 1 FROM display_columns WHERE column_name='tomb'").fetchone():
            conn.execute(
                "INSERT INTO display_columns VALUES (?, ?, ?, ?, ?, ?, ?)",
                ("tomb", "tomb", "hidden", columns.index("tomb"), 0, 0, None),
            )
        column_sql = ", ".join(f'"{name}"' for name in columns)
        placeholders = ", ".join("?" for _ in columns)
        conn.executemany(
            f"INSERT INTO database_info ({column_sql}) VALUES ({placeholders})",
            [[row[name] for name in columns] for row in pending],
        )
        total, tomb = conn.execute("SELECT COUNT(*), SUM(tomb) FROM database_info").fetchone()
        if total != len(workbook_rows) + len(pending) or tomb != 95:
            raise ValueError("Unexpected collection counts; import cancelled.")
        os.replace(staged, EXCEL_PATH)
        replaced_workbook = True
        conn.commit()
        committed = True
        report = {
            "added": len(pending), "total": total, "tomb": tomb,
            "source": str(SOURCE_PATH), "classification_source": str(CLASSIFICATION_PATH),
            "excel_before_sha256": hashlib.sha256(excel_before).hexdigest(),
            "backup": str(backup),
            "records": [{"id": row["id"], "database_name": row["database_name"]} for row in records],
        }
        (backup / "import_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
        print(json.dumps({key: value for key, value in report.items() if key != "records"}))
    except Exception:
        conn.rollback()
        if replaced_workbook and not committed:
            shutil.copy2(backup / "data.xlsx", EXCEL_PATH)
        raise
    finally:
        conn.close()
        book.close()


if __name__ == "__main__":
    main()
